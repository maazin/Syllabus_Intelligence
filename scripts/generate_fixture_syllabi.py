#!/usr/bin/env python3
"""Generate the sample syllabi in `tests/fixtures/syllabi/` — PRD section 26.

These are synthetic documents written to exercise specific hard cases from the
spec, not real syllabi. They exist so the full pipeline runs end to end on day
one without a live LLM budget or institutional data. They are explicitly NOT
the golden set: section 16's eval gate needs 100 hand-labeled *real* syllabi
before it means anything.

Coverage, matching section 26's table:
  - native-text PDF with a grading table
  - DOCX with a grading table
  - "Week N" style dates (no specific days)
  - a weekday/date mismatch (9.2's most common real-world defect)
  - a syllabus deferring the final to the registrar matrix (9.5)
  - a scanned-style PDF (image-only, no text layer) to exercise OCR routing
"""

from __future__ import annotations

import io
from pathlib import Path

FIXTURES = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "syllabi"


CS_LECTURE = {
    "title": "COP 4530 - Data Structures, Algorithms and Generic Programming",
    "meta": [
        "Section 003 | Fall 2026 | 3 credits",
        "Instructor: Dr. Alice Nakamura (a.nakamura@example.edu)",
        "Meeting pattern: MWF 10:00-10:50, Love Building 301",
        "Office hours: Wednesdays 1:00-3:00pm",
    ],
    "grade_table": [
        ["Category", "Weight", "Count", "Drop Lowest"],
        ["Exams", "40%", "2", "0"],
        ["Problem Sets", "20%", "10", "2"],
        ["Programming Projects", "25%", "3", "0"],
        ["Final Exam", "15%", "1", "0"],
    ],
    "schedule_table": [
        ["Date", "Topic", "Due"],
        ["Sep 4", "Asymptotic analysis", "Problem Set 1"],
        ["Sep 18", "Linked lists", "Problem Set 2"],
        ["Oct 2", "Trees and traversal", "Project 1"],
        ["Tuesday, October 14", "Midterm Exam 1", "Midterm Exam 1"],
        ["Oct 30", "Hash tables", "Problem Set 5"],
        ["Nov 13", "Graph algorithms", "Project 2"],
        ["Nov 20", "Midterm Exam 2", "Midterm Exam 2"],
        ["Dec 4", "Review", "Project 3"],
    ],
    "policies": [
        "Attendance is not graded, but you are responsible for all material presented in class.",
        "Late work: assignments lose 10% per day late, up to a maximum of three days. "
        "After three days, no credit is given.",
        "Group work: Programming Projects 2 and 3 are completed in pairs.",
        "Final exam: cumulative. See the university final exam schedule for the date and time.",
        "AI policy: Use of generative AI tools is permitted provided you disclose "
        "such use in a comment at the top of each submitted file.",
        "Required text: Weiss, Data Structures and Algorithm Analysis in C++, 4th ed. "
        "ISBN 978-0132847377. Approximately $120.",
    ],
}

HUMANITIES_SEMINAR = {
    "title": "ENG 3014 - Modern American Literature",
    "meta": [
        "Section 001 | Fall 2026 | 3 credits",
        "Instructor: Prof. Daniel Whitfield (d.whitfield@example.edu)",
        "Meeting pattern: TR 14:00-15:15, Williams Building 216",
    ],
    "grade_table": [
        ["Component", "Percentage"],
        ["Participation", "15%"],
        ["Reading Responses", "20%"],
        ["Short Paper", "25%"],
        ["Final Paper", "40%"],
    ],
    # Deliberately week-based: exercises the week_only precision path (9.3).
    "schedule_table": [
        ["Week", "Reading", "Due"],
        ["Week 2", "Fitzgerald", "Reading Response 1"],
        ["Week 5", "Hemingway", "Reading Response 2"],
        ["Week 7", "Faulkner", "Short Paper"],
        ["Week 10", "Morrison", "Reading Response 3"],
        ["Week 14", "Workshop", "Final Paper"],
    ],
    "policies": [
        "Attendance is graded and counts toward the participation component. "
        "More than three unexcused absences will lower your final grade.",
        "Late work: papers submitted after the deadline are not accepted except "
        "in documented emergencies.",
        "There is no final exam in this course. The Final Paper serves as the "
        "culminating assessment.",
        "AI policy: The use of generative AI tools for any graded work is prohibited.",
        "Reading responses are due Sundays at 11:59pm via the course site.",
    ],
}


def _pdf(path: Path, spec: dict) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story = [Paragraph(spec["title"], styles["Title"]), Spacer(1, 10)]

    for line in spec["meta"]:
        story.append(Paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 14))

    story.append(Paragraph("Grading", styles["Heading2"]))
    story.append(_table(spec["grade_table"], Table, TableStyle, colors))
    story.append(Spacer(1, 14))

    story.append(Paragraph("Course Schedule", styles["Heading2"]))
    story.append(_table(spec["schedule_table"], Table, TableStyle, colors))
    story.append(Spacer(1, 14))

    story.append(Paragraph("Course Policies", styles["Heading2"]))
    for policy in spec["policies"]:
        story.append(Paragraph(policy, styles["Normal"]))
        story.append(Spacer(1, 6))

    SimpleDocTemplate(str(path), pagesize=LETTER).build(story)


def _table(rows, Table, TableStyle, colors):
    table = Table(rows, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return table


def _docx(path: Path, spec: dict) -> None:
    import docx

    document = docx.Document()
    document.add_heading(spec["title"], level=0)
    for line in spec["meta"]:
        document.add_paragraph(line)

    document.add_heading("Grading", level=1)
    _docx_table(document, spec["grade_table"])

    document.add_heading("Course Schedule", level=1)
    _docx_table(document, spec["schedule_table"])

    document.add_heading("Course Policies", level=1)
    for policy in spec["policies"]:
        document.add_paragraph(policy)

    document.save(str(path))


def _docx_table(document, rows) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            table.cell(r, c).text = cell


def _scanned_pdf(path: Path, spec: dict) -> None:
    """A PDF whose pages are images — no text layer at all.

    Exercises the section 7 detection heuristic (under ~100 chars per page
    routes to OCR). Roughly 10-15% of real syllabi arrive like this.
    """
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdfcanvas

    try:
        from PIL import Image, ImageDraw
    except ImportError:
        print("  ! Pillow not installed; skipping the scanned fixture")
        return

    lines = [spec["title"], ""]
    lines.extend(spec["meta"])
    lines.append("")
    for row in spec["grade_table"]:
        lines.append("   ".join(row))

    image = Image.new("RGB", (1275, 1650), "white")
    draw = ImageDraw.Draw(image)
    y = 80
    for line in lines:
        draw.text((90, y), line, fill="black")
        y += 26

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    pdf = pdfcanvas.Canvas(str(path), pagesize=LETTER)
    pdf.drawImage(ImageReader(buffer), 0, 0, width=LETTER[0], height=LETTER[1])
    pdf.showPage()
    pdf.save()


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)

    targets = [
        ("cs_lecture_native_pdf.pdf", CS_LECTURE, _pdf),
        ("humanities_seminar_week_dates.pdf", HUMANITIES_SEMINAR, _pdf),
        ("cs_lecture_grading_table.docx", CS_LECTURE, _docx),
        ("humanities_seminar.docx", HUMANITIES_SEMINAR, _docx),
        ("cs_lecture_scanned.pdf", CS_LECTURE, _scanned_pdf),
    ]

    for filename, spec, writer in targets:
        path = FIXTURES / filename
        writer(path, spec)
        if path.exists():
            print(
                f"  wrote {path.relative_to(FIXTURES.parents[2])} ({path.stat().st_size:,} bytes)"
            )

    plain = FIXTURES / "cs_lecture_pasted.txt"
    body = [CS_LECTURE["title"], ""]
    body.extend(CS_LECTURE["meta"])
    body.append("")
    body.append("Grading:")
    for row in CS_LECTURE["grade_table"][1:]:
        body.append(f"  {row[0]}: {row[1]}")
    body.append("")
    body.append("Schedule:")
    for row in CS_LECTURE["schedule_table"][1:]:
        body.append(f"  {row[0]} - {row[1]} - due: {row[2]}")
    body.append("")
    body.extend(CS_LECTURE["policies"])
    plain.write_text("\n".join(body), encoding="utf-8")
    print(f"  wrote {plain.relative_to(FIXTURES.parents[2])} ({plain.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
