"""Ingestion pipeline — PRD section 7.

    upload -> virus scan -> dedupe (sha256) -> format detect
       -> text layer extraction -> layout-aware normalization -> chunk

Two implementation notes from the PRD drive the shape of this module:

1. **Tables matter more than prose.** Grading breakdowns and course schedules are
   almost always tables, and a text extractor that flattens a table into word soup
   destroys weight extraction. Tables are preserved as markdown before the model
   sees anything, and `extract_tables_as_markdown` is evaluated independently of
   extraction quality (Appendix B step 6).
2. **Scanned detection is a character-count heuristic.** Under ~100 chars per page
   means image-only; roughly 10-15% of real syllabi arrive that way.
"""

from __future__ import annotations

import hashlib
import logging
import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path

logger = logging.getLogger(__name__)

#: Under this many characters per page, the page has no usable text layer (section 7).
SCANNED_CHARS_PER_PAGE_THRESHOLD = 100

#: Chunk above this token estimate (section 7). Syllabi are 2-8k tokens, so most
#: never chunk; the long ones are usually appendix-heavy.
CHUNK_TOKEN_THRESHOLD = 12_000

#: Rough chars-per-token for English prose. Only used to decide *whether* to chunk,
#: never to bill or budget, so an approximation is fine.
CHARS_PER_TOKEN = 4


class DocumentFormat(StrEnum):
    PDF_NATIVE = "pdf_native"
    PDF_SCANNED = "pdf_scanned"
    DOCX = "docx"
    IMAGE = "image"
    HTML = "html"
    TEXT = "text"
    UNKNOWN = "unknown"


class IngestionError(Exception):
    """Raised when a document cannot be turned into text at all.

    US-1 requires one file's failure not to block the other four, so callers
    catch this per-document and mark only that document failed.
    """


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    tables_markdown: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.text.strip())


@dataclass
class IngestedDocument:
    """The output of ingestion: text the extraction prompts can consume."""

    sha256: str
    format: DocumentFormat
    pages: list[ExtractedPage]
    is_scanned: bool
    ocr_applied: bool = False

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def full_text(self) -> str:
        """Pages joined with explicit page markers.

        The markers are load-bearing: `page_ref` in the extraction schema (8.2)
        is what lets the review screen show "page 3, these two lines" (US-3),
        and the model can only report a page it can see.
        """
        parts: list[str] = []
        for page in self.pages:
            parts.append(f"[page {page.page_number}]")
            if page.text.strip():
                parts.append(page.text.strip())
            for table in page.tables_markdown:
                parts.append(table)
        return "\n\n".join(parts)

    @property
    def estimated_tokens(self) -> int:
        return len(self.full_text) // CHARS_PER_TOKEN

    @property
    def total_chars(self) -> int:
        return sum(p.char_count for p in self.pages)


def content_hash(data: bytes) -> str:
    """SHA-256 of the raw upload, for the dedup step in section 7.

    Dedup is load-bearing rather than an optimization (18.3): dozens of students
    in one section upload the identical file, and parsing once per section is
    what keeps token spend and Neon compute-hours inside the free tiers during
    week 0. It also means later uploaders get instant results.
    """
    return hashlib.sha256(data).hexdigest()


def detect_format(filename: str, data: bytes) -> DocumentFormat:
    """Format detection by magic bytes first, extension second.

    Magic bytes win because browsers and phones mislabel uploads constantly —
    a HEIC photo arriving as `.jpg` is routine.
    """
    if data.startswith(b"%PDF"):
        return DocumentFormat.PDF_NATIVE  # refined to PDF_SCANNED after text extraction
    if data.startswith(b"PK\x03\x04") and filename.lower().endswith(".docx"):
        return DocumentFormat.DOCX
    if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n":
        return DocumentFormat.IMAGE
    if len(data) > 12 and data[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1"):
        return DocumentFormat.IMAGE

    suffix = Path(filename).suffix.lower()
    if suffix in (".html", ".htm"):
        return DocumentFormat.HTML
    if suffix in (".txt", ".md"):
        return DocumentFormat.TEXT
    if suffix in (".png", ".jpg", ".jpeg", ".heic", ".webp"):
        return DocumentFormat.IMAGE
    if suffix == ".docx":
        return DocumentFormat.DOCX

    # A pasted-text upload has no filename and no magic bytes. Decodability alone
    # is too weak a test — control bytes are valid UTF-8 — so require the sample to
    # be mostly printable before calling it text.
    try:
        sample = data[:2048].decode("utf-8")
    except UnicodeDecodeError:
        return DocumentFormat.UNKNOWN
    if not sample.strip():
        return DocumentFormat.UNKNOWN

    printable = sum(1 for ch in sample if ch.isprintable() or ch in "\n\r\t")
    return DocumentFormat.TEXT if printable / len(sample) > 0.9 else DocumentFormat.UNKNOWN


def table_to_markdown(rows: Sequence[Sequence[str | None]]) -> str:
    """Render a extracted table as a markdown table.

    Markdown, specifically, because it survives tokenization with the row/column
    relationship intact — which is the entire reason section 7 calls out tables
    as mattering more than prose.
    """
    cleaned = [
        [(cell or "").strip().replace("\n", " ").replace("|", "\\|") for cell in row]
        for row in rows
        if row and any((cell or "").strip() for cell in row)
    ]
    if not cleaned:
        return ""

    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]

    header, *body = cleaned
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> tuple[list[ExtractedPage], bool]:
    """Native PDF text + table structure via pdfplumber.

    Returns (pages, looks_scanned). A PDF whose text layer is essentially empty
    is a scan and must be routed to OCR rather than handed to the model as
    thirty blank pages.
    """
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency presence is environmental
        raise IngestionError(
            "pdfplumber is required for PDF ingestion; install the worker dependencies"
        ) from exc

    import io

    pages: list[ExtractedPage] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                tables: list[str] = []
                for raw_table in page.extract_tables() or []:
                    markdown = table_to_markdown(raw_table)
                    if markdown:
                        tables.append(markdown)
                pages.append(ExtractedPage(page_number=index, text=text, tables_markdown=tables))
    except Exception as exc:
        raise IngestionError(f"Could not read PDF: {exc}") from exc

    if not pages:
        raise IngestionError("PDF contained no pages")

    avg_chars = sum(p.char_count for p in pages) / len(pages)
    return pages, avg_chars < SCANNED_CHARS_PER_PAGE_THRESHOLD


def _extract_docx(data: bytes) -> list[ExtractedPage]:
    """DOCX text and tables. DOCX has no page concept, so everything is page 1.

    Grading breakdowns in DOCX are nearly always real Word tables, so the table
    pass here matters as much as it does for PDF.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise IngestionError(
            "python-docx is required for DOCX ingestion; install the worker dependencies"
        ) from exc

    import io

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"Could not read DOCX: {exc}") from exc

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables = []
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown = table_to_markdown(rows)
        if markdown:
            tables.append(markdown)

    return [ExtractedPage(page_number=1, text="\n".join(paragraphs), tables_markdown=tables)]


def _extract_html(data: bytes) -> list[ExtractedPage]:
    """Strip an LMS HTML export down to readable text, keeping tables as markdown."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover
        raise IngestionError(
            "beautifulsoup4 is required for HTML ingestion; install the worker dependencies"
        ) from exc

    soup = BeautifulSoup(data.decode("utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()

    tables: list[str] = []
    for table in soup.find_all("table"):
        rows = [
            [cell.get_text(strip=True) for cell in row.find_all(["td", "th"])]
            for row in table.find_all("tr")
        ]
        markdown = table_to_markdown(rows)
        if markdown:
            tables.append(markdown)
        table.decompose()  # remove so its text is not duplicated into the prose

    text = re.sub(r"\n{3,}", "\n\n", soup.get_text("\n", strip=True))
    return [ExtractedPage(page_number=1, text=text, tables_markdown=tables)]


def _ocr(data: bytes, mode: str) -> list[ExtractedPage]:
    """Route a scanned document to the configured OCR engine (`OCR_MODE`, section 24).

    Docling is the default per 18.1: it handles table *structure*, which Tesseract
    has no API for at all, and tables are the deciding requirement. It is imported
    lazily because the model weights are large and most documents never need them.
    """
    if mode == "docling":
        try:
            from services.worker.tasks.ocr_docling import ocr_with_docling
        except ImportError:
            # Unresolvable to a type checker on purpose; see extract.py.
            from ocr_docling import ocr_with_docling  # type: ignore[no-redef,import-not-found]
        return ocr_with_docling(data)

    raise IngestionError(f"Unsupported OCR_MODE {mode!r}")


def ingest(
    filename: str,
    data: bytes,
    *,
    ocr_mode: str = "docling",
    allow_ocr: bool = True,
) -> IngestedDocument:
    """Run the full section 7 pipeline over one uploaded file.

    Raises `IngestionError` on anything unreadable; the caller marks that one
    document failed and leaves the rest of the batch alone (US-1).
    """
    sha = content_hash(data)
    fmt = detect_format(filename, data)
    ocr_applied = False

    if fmt is DocumentFormat.UNKNOWN:
        raise IngestionError(f"Unsupported file type for {filename!r}")

    if fmt is DocumentFormat.PDF_NATIVE:
        pages, looks_scanned = _extract_pdf(data)
        if looks_scanned:
            fmt = DocumentFormat.PDF_SCANNED
            if allow_ocr:
                logger.info("Routing %s to OCR: text layer below threshold", filename)
                pages = _ocr(data, ocr_mode)
                ocr_applied = True
    elif fmt is DocumentFormat.DOCX:
        pages = _extract_docx(data)
    elif fmt is DocumentFormat.HTML:
        pages = _extract_html(data)
    elif fmt is DocumentFormat.TEXT:
        pages = [ExtractedPage(page_number=1, text=data.decode("utf-8", errors="replace"))]
    elif fmt is DocumentFormat.IMAGE:
        if not allow_ocr:
            raise IngestionError("Image upload requires OCR, which is disabled")
        pages = _ocr(data, ocr_mode)
        ocr_applied = True
    else:  # pragma: no cover - every enum member is handled above
        raise IngestionError(f"Unhandled format {fmt}")

    document = IngestedDocument(
        sha256=sha,
        format=fmt,
        pages=pages,
        is_scanned=fmt is DocumentFormat.PDF_SCANNED or ocr_applied,
        ocr_applied=ocr_applied,
    )

    if document.total_chars == 0:
        raise IngestionError(
            f"No text could be extracted from {filename!r}. "
            "Route to manual entry rather than showing an empty result."
        )
    return document


def chunk_text(text: str, max_tokens: int = CHUNK_TOKEN_THRESHOLD) -> list[str]:
    """Split oversized documents on page boundaries (section 7).

    Page boundaries specifically, so a grading table is never cut in half — a
    split mid-table would reintroduce exactly the word-soup failure the table
    handling exists to prevent.
    """
    if len(text) // CHARS_PER_TOKEN <= max_tokens:
        return [text]

    max_chars = max_tokens * CHARS_PER_TOKEN
    page_blocks = re.split(r"(?=\[page \d+\])", text)

    chunks: list[str] = []
    current = ""
    for block in page_blocks:
        if not block:
            continue
        if current and len(current) + len(block) > max_chars:
            chunks.append(current.strip())
            current = block
        else:
            current += block
    if current.strip():
        chunks.append(current.strip())
    return chunks or [text[:max_chars]]
